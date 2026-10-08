"""POS Billing & Inventory Management - FastAPI backend."""
import os
import logging
import uuid
import asyncio
from pathlib import Path
from datetime import datetime, timezone, timedelta, date
from typing import List, Optional, Literal

import bcrypt
import jwt
from fastapi import FastAPI, APIRouter, Depends, HTTPException, status, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr, ConfigDict

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

# ---------------- DB ----------------
mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

# ---------------- App ----------------
app = FastAPI(title="POS Billing & Inventory API")
api = APIRouter(prefix="/api")
security = HTTPBearer(auto_error=False)

JWT_SECRET = os.environ.get("JWT_SECRET", "pos-secret-change-me-please-3782")
JWT_ALGO = "HS256"
JWT_EXP_HOURS = 12

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("pos")


# ---------------- Helpers ----------------
def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def hash_pw(pw: str) -> str:
    return bcrypt.hashpw(pw.encode(), bcrypt.gensalt()).decode()


def verify_pw(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode(), hashed.encode())
    except Exception:
        return False


def make_token(user: dict) -> str:
    payload = {
        "sub": user["id"],
        "username": user["username"],
        "role": user["role"],
        "exp": datetime.now(timezone.utc) + timedelta(hours=JWT_EXP_HOURS),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGO)


async def get_current_user(creds: Optional[HTTPAuthorizationCredentials] = Depends(security)) -> dict:
    if not creds or not creds.credentials:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALGO])
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "passwordHash": 0})
    if not user:
        raise HTTPException(401, "User not found")
    if user.get("status") != "active":
        raise HTTPException(403, "User is inactive")
    return user


async def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user["role"] != "admin":
        raise HTTPException(403, "Admin access required")
    return user


def strip_id(doc: Optional[dict]) -> Optional[dict]:
    if not doc:
        return None
    doc.pop("_id", None)
    return doc


# ---------------- Pydantic Models ----------------
class LoginRequest(BaseModel):
    username: str
    password: str


class UserCreate(BaseModel):
    username: str
    fullName: str
    mobileNumber: Optional[str] = ""
    password: str
    role: Literal["admin", "cashier"] = "cashier"


class UserUpdate(BaseModel):
    fullName: Optional[str] = None
    mobileNumber: Optional[str] = None
    status: Optional[Literal["active", "inactive"]] = None
    role: Optional[Literal["admin", "cashier"]] = None


class PasswordReset(BaseModel):
    newPassword: str


class CategoryIn(BaseModel):
    categoryName: str
    description: Optional[str] = ""
    status: Literal["active", "inactive"] = "active"


class SupplierIn(BaseModel):
    supplierName: str
    contactPerson: Optional[str] = ""
    mobileNumber: Optional[str] = ""
    email: Optional[str] = ""
    address: Optional[str] = ""
    status: Literal["active", "inactive"] = "active"


class ProductIn(BaseModel):
    productName: str
    barcode: Optional[str] = ""
    sku: Optional[str] = ""
    category: Optional[str] = ""
    brand: Optional[str] = ""
    costPrice: float = 0
    sellingPrice: float = 0
    taxPercentage: float = 0
    stockQuantity: float = 0
    minimumStockLevel: float = 0
    supplier: Optional[str] = ""
    productImageUrl: Optional[str] = ""
    status: Literal["active", "inactive"] = "active"


class StockMovementIn(BaseModel):
    productId: str
    movementType: Literal["inward", "adjustment", "damage", "lost", "correction"]
    quantity: float
    reason: Optional[str] = ""


class CartItem(BaseModel):
    productId: str
    quantity: float
    discountAmount: float = 0


class SaleCreate(BaseModel):
    customerName: Optional[str] = ""
    items: List[CartItem]
    discountAmount: float = 0
    paymentMethod: Literal["cash", "card", "upi", "mixed"]
    payments: Optional[List[dict]] = None  # [{paymentMethod, amount}]


class StoreSettingsIn(BaseModel):
    storeName: str
    storeAddress: Optional[str] = ""
    mobileNumber: Optional[str] = ""
    email: Optional[str] = ""
    gstNumber: Optional[str] = ""
    taxEnabled: bool = True
    invoicePrefix: str = "INV"
    receiptFooter: Optional[str] = ""


# ---------------- Auth Routes ----------------
@api.post("/auth/login")
async def login(req: LoginRequest):
    user = await db.users.find_one({"username": req.username})
    if not user or not verify_pw(req.password, user["passwordHash"]):
        raise HTTPException(401, "Invalid username or password")
    if user.get("status") != "active":
        raise HTTPException(403, "Account is inactive")
    token = make_token(user)
    return {
        "token": token,
        "user": {
            "id": user["id"],
            "username": user["username"],
            "fullName": user["fullName"],
            "role": user["role"],
        },
    }


@api.post("/auth/logout")
async def logout(user: dict = Depends(get_current_user)):
    return {"success": True}


@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


# ---------------- User Management (admin) ----------------
@api.get("/users")
async def list_users(_: dict = Depends(require_admin)):
    cur = db.users.find({}, {"_id": 0, "passwordHash": 0}).sort("createdAt", -1)
    return await cur.to_list(1000)


@api.post("/users")
async def create_user(payload: UserCreate, _: dict = Depends(require_admin)):
    existing = await db.users.find_one({"username": payload.username})
    if existing:
        raise HTTPException(400, "Username already exists")
    doc = {
        "id": str(uuid.uuid4()),
        "username": payload.username,
        "fullName": payload.fullName,
        "mobileNumber": payload.mobileNumber or "",
        "passwordHash": hash_pw(payload.password),
        "role": payload.role,
        "status": "active",
        "createdAt": now_iso(),
        "updatedAt": now_iso(),
    }
    await db.users.insert_one(doc)
    doc.pop("passwordHash", None)
    doc.pop("_id", None)
    return doc


@api.put("/users/{user_id}")
async def update_user(user_id: str, payload: UserUpdate, _: dict = Depends(require_admin)):
    updates = {k: v for k, v in payload.model_dump().items() if v is not None}
    if not updates:
        raise HTTPException(400, "Nothing to update")
    updates["updatedAt"] = now_iso()
    res = await db.users.update_one({"id": user_id}, {"$set": updates})
    if res.matched_count == 0:
        raise HTTPException(404, "User not found")
    return {"success": True}


@api.post("/users/{user_id}/reset-password")
async def reset_password(user_id: str, payload: PasswordReset, _: dict = Depends(require_admin)):
    res = await db.users.update_one(
        {"id": user_id},
        {"$set": {"passwordHash": hash_pw(payload.newPassword), "updatedAt": now_iso()}},
    )
    if res.matched_count == 0:
        raise HTTPException(404, "User not found")
    return {"success": True}


# ---------------- Categories ----------------
@api.get("/categories")
async def list_categories(user: dict = Depends(get_current_user)):
    cur = db.categories.find({}, {"_id": 0}).sort("categoryName", 1)
    return await cur.to_list(1000)


@api.post("/categories")
async def create_category(payload: CategoryIn, _: dict = Depends(require_admin)):
    if await db.categories.find_one({"categoryName": payload.categoryName}):
        raise HTTPException(400, "Category name already exists")
    doc = {
        "id": str(uuid.uuid4()),
        **payload.model_dump(),
        "createdAt": now_iso(),
        "updatedAt": now_iso(),
    }
    await db.categories.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.put("/categories/{cid}")
async def update_category(cid: str, payload: CategoryIn, _: dict = Depends(require_admin)):
    res = await db.categories.update_one(
        {"id": cid}, {"$set": {**payload.model_dump(), "updatedAt": now_iso()}}
    )
    if res.matched_count == 0:
        raise HTTPException(404, "Category not found")
    return {"success": True}


@api.delete("/categories/{cid}")
async def delete_category(cid: str, _: dict = Depends(require_admin)):
    in_use = await db.products.find_one({"category": cid})
    if in_use:
        raise HTTPException(400, "Category in use by products")
    await db.categories.delete_one({"id": cid})
    return {"success": True}


# ---------------- Suppliers ----------------
@api.get("/suppliers")
async def list_suppliers(_: dict = Depends(require_admin)):
    cur = db.suppliers.find({}, {"_id": 0}).sort("supplierName", 1)
    return await cur.to_list(1000)


@api.post("/suppliers")
async def create_supplier(payload: SupplierIn, _: dict = Depends(require_admin)):
    doc = {
        "id": str(uuid.uuid4()),
        **payload.model_dump(),
        "createdAt": now_iso(),
        "updatedAt": now_iso(),
    }
    await db.suppliers.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.put("/suppliers/{sid}")
async def update_supplier(sid: str, payload: SupplierIn, _: dict = Depends(require_admin)):
    res = await db.suppliers.update_one(
        {"id": sid}, {"$set": {**payload.model_dump(), "updatedAt": now_iso()}}
    )
    if res.matched_count == 0:
        raise HTTPException(404, "Supplier not found")
    return {"success": True}


@api.delete("/suppliers/{sid}")
async def delete_supplier(sid: str, _: dict = Depends(require_admin)):
    await db.suppliers.delete_one({"id": sid})
    return {"success": True}


# ---------------- Products ----------------
@api.get("/products")
async def list_products(
    user: dict = Depends(get_current_user),
    q: Optional[str] = None,
    activeOnly: bool = False,
    lowStock: bool = False,
    limit: int = 200,
):
    query: dict = {}
    if activeOnly or user["role"] == "cashier":
        query["status"] = "active"
    if q:
        query["$or"] = [
            {"productName": {"$regex": q, "$options": "i"}},
            {"barcode": q},
            {"sku": {"$regex": q, "$options": "i"}},
            {"brand": {"$regex": q, "$options": "i"}},
        ]
    cur = db.products.find(query, {"_id": 0}).sort("productName", 1).limit(limit)
    items = await cur.to_list(limit)
    if lowStock:
        items = [p for p in items if p.get("stockQuantity", 0) <= p.get("minimumStockLevel", 0)]
    return items


@api.get("/products/{pid}")
async def get_product(pid: str, _: dict = Depends(get_current_user)):
    doc = await db.products.find_one({"id": pid}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Product not found")
    return doc


@api.post("/products")
async def create_product(payload: ProductIn, _: dict = Depends(require_admin)):
    data = payload.model_dump()
    if data.get("barcode"):
        if await db.products.find_one({"barcode": data["barcode"]}):
            raise HTTPException(400, "Barcode already in use")
    if data.get("sku"):
        if await db.products.find_one({"sku": data["sku"]}):
            raise HTTPException(400, "SKU already in use")
    doc = {
        "id": str(uuid.uuid4()),
        **data,
        "createdAt": now_iso(),
        "updatedAt": now_iso(),
    }
    await db.products.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api.put("/products/{pid}")
async def update_product(pid: str, payload: ProductIn, _: dict = Depends(require_admin)):
    data = payload.model_dump()
    res = await db.products.update_one(
        {"id": pid}, {"$set": {**data, "updatedAt": now_iso()}}
    )
    if res.matched_count == 0:
        raise HTTPException(404, "Product not found")
    return {"success": True}


@api.delete("/products/{pid}")
async def delete_product(pid: str, _: dict = Depends(require_admin)):
    await db.products.delete_one({"id": pid})
    return {"success": True}


@api.get("/products-template")
async def products_template(_: dict = Depends(require_admin)):
    from openpyxl import Workbook
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    headers = [
        "productName", "barcode", "sku", "category", "brand",
        "costPrice", "sellingPrice", "taxPercentage",
        "stockQuantity", "minimumStockLevel", "supplier", "status",
    ]
    ws.append(headers)
    ws.append([
        "Example Product", "8900000000001", "SKU-001", "Beverages", "Acme",
        100, 150, 5, 50, 10, "Global Distributors", "active",
    ])
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=products_template.xlsx"},
    )


@api.post("/products/bulk-upload")
async def bulk_upload_products(file: UploadFile = File(...), _: dict = Depends(require_admin)):
    from openpyxl import load_workbook
    import io

    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Please upload a .xlsx file")
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Could not read file: {e}")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise HTTPException(400, "File is empty")

    headers = [str(h).strip() if h else "" for h in rows[0]]
    expected = {
        "productName", "barcode", "sku", "category", "brand",
        "costPrice", "sellingPrice", "taxPercentage",
        "stockQuantity", "minimumStockLevel", "supplier", "status",
    }
    if not set(headers) & expected:
        raise HTTPException(400, "Header row does not match expected fields")

    # Build lookups
    cats = await db.categories.find({}, {"_id": 0, "id": 1, "categoryName": 1}).to_list(1000)
    cat_lookup = {c["categoryName"].lower(): c["id"] for c in cats}
    cat_by_id = {c["id"]: c["id"] for c in cats}
    supps = await db.suppliers.find({}, {"_id": 0, "id": 1, "supplierName": 1}).to_list(1000)
    supp_lookup = {s["supplierName"].lower(): s["id"] for s in supps}
    supp_by_id = {s["id"]: s["id"] for s in supps}

    created, updated, errors = 0, 0, []
    for idx, row in enumerate(rows[1:], start=2):
        record = {h: row[i] for i, h in enumerate(headers) if h}
        name = (record.get("productName") or "").strip() if isinstance(record.get("productName"), str) else record.get("productName")
        if not name:
            errors.append({"row": idx, "error": "Missing productName"})
            continue

        def num(v, default=0):
            try:
                return float(v) if v not in (None, "") else default
            except (ValueError, TypeError):
                return default

        cat_raw = (record.get("category") or "")
        cat_id = cat_by_id.get(str(cat_raw)) or cat_lookup.get(str(cat_raw).lower(), "") if cat_raw else ""
        supp_raw = (record.get("supplier") or "")
        supp_id = supp_by_id.get(str(supp_raw)) or supp_lookup.get(str(supp_raw).lower(), "") if supp_raw else ""

        data = {
            "productName": str(name).strip(),
            "barcode": str(record.get("barcode") or "").strip(),
            "sku": str(record.get("sku") or "").strip(),
            "category": cat_id,
            "brand": str(record.get("brand") or "").strip(),
            "costPrice": num(record.get("costPrice")),
            "sellingPrice": num(record.get("sellingPrice")),
            "taxPercentage": num(record.get("taxPercentage")),
            "stockQuantity": num(record.get("stockQuantity")),
            "minimumStockLevel": num(record.get("minimumStockLevel")),
            "supplier": supp_id,
            "productImageUrl": "",
            "status": str(record.get("status") or "active").strip().lower(),
        }
        if data["status"] not in ("active", "inactive"):
            data["status"] = "active"

        existing = None
        if data["barcode"]:
            existing = await db.products.find_one({"barcode": data["barcode"]})
        if not existing and data["sku"]:
            existing = await db.products.find_one({"sku": data["sku"]})
        if not existing:
            existing = await db.products.find_one({"productName": data["productName"]})

        try:
            if existing:
                await db.products.update_one(
                    {"id": existing["id"]},
                    {"$set": {**data, "updatedAt": now_iso()}},
                )
                updated += 1
            else:
                await db.products.insert_one({
                    "id": str(uuid.uuid4()),
                    **data,
                    "createdAt": now_iso(),
                    "updatedAt": now_iso(),
                })
                created += 1
        except Exception as e:
            errors.append({"row": idx, "error": str(e)})

    return {"created": created, "updated": updated, "errors": errors, "totalRows": len(rows) - 1}


# ---------------- Stock Movements ----------------
@api.post("/stock-movements")
async def create_stock_movement(payload: StockMovementIn, admin: dict = Depends(require_admin)):
    product = await db.products.find_one({"id": payload.productId})
    if not product:
        raise HTTPException(404, "Product not found")
    prev = float(product.get("stockQuantity", 0))
    qty = float(payload.quantity)
    if payload.movementType in ("inward", "correction"):
        new_stock = prev + qty
    elif payload.movementType == "adjustment":
        new_stock = qty  # set to value
    else:  # damage, lost
        new_stock = prev - qty
    if new_stock < 0:
        raise HTTPException(400, "Resulting stock would be negative")
    mv = {
        "id": str(uuid.uuid4()),
        "product": payload.productId,
        "productName": product["productName"],
        "movementType": payload.movementType,
        "quantity": qty,
        "previousStock": prev,
        "newStock": new_stock,
        "reason": payload.reason or "",
        "createdBy": admin["id"],
        "createdByName": admin["fullName"],
        "createdAt": now_iso(),
    }
    await db.stock_movements.insert_one(mv)
    await db.products.update_one(
        {"id": payload.productId},
        {"$set": {"stockQuantity": new_stock, "updatedAt": now_iso()}},
    )
    mv.pop("_id", None)
    return mv


@api.get("/stock-movements")
async def list_stock_movements(
    _: dict = Depends(require_admin),
    productId: Optional[str] = None,
    limit: int = 200,
):
    q = {"product": productId} if productId else {}
    cur = db.stock_movements.find(q, {"_id": 0}).sort("createdAt", -1).limit(limit)
    return await cur.to_list(limit)


# ---------------- Billing ----------------
async def _next_invoice_number() -> str:
    settings = await db.store_settings.find_one({"id": "default"}, {"_id": 0})
    prefix = (settings or {}).get("invoicePrefix", "INV")
    # use date + counter
    today = datetime.now(timezone.utc).strftime("%Y%m%d")
    counter_doc = await db.counters.find_one_and_update(
        {"id": f"inv-{today}"},
        {"$inc": {"seq": 1}},
        upsert=True,
        return_document=True,
    )
    seq = (counter_doc or {}).get("seq", 1)
    return f"{prefix}-{today}-{seq:04d}"


@api.post("/sales")
async def create_sale(payload: SaleCreate, user: dict = Depends(get_current_user)):
    if not payload.items:
        raise HTTPException(400, "Cart is empty")

    # Fetch all products
    product_ids = [it.productId for it in payload.items]
    products_cursor = db.products.find({"id": {"$in": product_ids}}, {"_id": 0})
    products = {p["id"]: p async for p in products_cursor}

    settings = await db.store_settings.find_one({"id": "default"}, {"_id": 0})
    tax_enabled = (settings or {}).get("taxEnabled", True)

    sale_items: List[dict] = []
    subtotal = 0.0
    tax_total = 0.0
    for it in payload.items:
        p = products.get(it.productId)
        if not p:
            raise HTTPException(400, f"Product {it.productId} not found")
        if p.get("status") != "active":
            raise HTTPException(400, f"Product {p['productName']} is inactive")
        if float(p.get("stockQuantity", 0)) < it.quantity:
            raise HTTPException(400, f"Insufficient stock for {p['productName']}")
        unit_price = float(p["sellingPrice"])
        line_subtotal = unit_price * it.quantity - it.discountAmount
        if line_subtotal < 0:
            line_subtotal = 0
        tax_pct = float(p.get("taxPercentage", 0)) if tax_enabled else 0
        line_tax = round(line_subtotal * tax_pct / 100.0, 2)
        line_total = round(line_subtotal + line_tax, 2)
        sale_items.append({
            "id": str(uuid.uuid4()),
            "product": p["id"],
            "productName": p["productName"],
            "barcode": p.get("barcode", ""),
            "quantity": it.quantity,
            "unitPrice": unit_price,
            "discountAmount": it.discountAmount,
            "taxPercentage": tax_pct,
            "taxAmount": line_tax,
            "totalAmount": line_total,
        })
        subtotal += unit_price * it.quantity
        tax_total += line_tax

    discount_amount = float(payload.discountAmount or 0)
    grand_total = round(subtotal - discount_amount + tax_total - sum(i["discountAmount"] for i in sale_items), 2)
    if grand_total < 0:
        grand_total = 0
    invoice_no = await _next_invoice_number()

    sale_doc = {
        "id": str(uuid.uuid4()),
        "invoiceNumber": invoice_no,
        "cashier": user["id"],
        "cashierName": user["fullName"],
        "customerName": payload.customerName or "",
        "saleDate": now_iso(),
        "subtotal": round(subtotal, 2),
        "discountAmount": discount_amount,
        "taxAmount": round(tax_total, 2),
        "grandTotal": grand_total,
        "paymentMethod": payload.paymentMethod,
        "status": "completed",
        "createdAt": now_iso(),
        "updatedAt": now_iso(),
    }
    await db.sales.insert_one(sale_doc)

    # Save sale items
    for si in sale_items:
        si["sale"] = sale_doc["id"]
        si["invoiceNumber"] = invoice_no
        si["createdAt"] = now_iso()
        await db.sale_items.insert_one(si)

    # Payments
    payments = payload.payments or [{"paymentMethod": payload.paymentMethod, "amount": grand_total}]
    for pay in payments:
        await db.payments.insert_one({
            "id": str(uuid.uuid4()),
            "sale": sale_doc["id"],
            "paymentMethod": pay.get("paymentMethod", payload.paymentMethod),
            "amount": float(pay.get("amount", 0)),
            "createdAt": now_iso(),
        })

    # Reduce stock and log movements
    for it in payload.items:
        p = products[it.productId]
        prev = float(p.get("stockQuantity", 0))
        new_stock = prev - it.quantity
        await db.products.update_one(
            {"id": p["id"]},
            {"$set": {"stockQuantity": new_stock, "updatedAt": now_iso()}},
        )
        await db.stock_movements.insert_one({
            "id": str(uuid.uuid4()),
            "product": p["id"],
            "productName": p["productName"],
            "movementType": "sale",
            "quantity": it.quantity,
            "previousStock": prev,
            "newStock": new_stock,
            "reason": f"Sale {invoice_no}",
            "createdBy": user["id"],
            "createdByName": user["fullName"],
            "createdAt": now_iso(),
        })

    # Build receipt
    sale_doc.pop("_id", None)
    for si in sale_items:
        si.pop("_id", None)
    return {"sale": sale_doc, "items": sale_items, "store": settings or {}}


@api.get("/sales")
async def list_sales(
    user: dict = Depends(get_current_user),
    mine: bool = False,
    fromDate: Optional[str] = None,
    toDate: Optional[str] = None,
    limit: int = 200,
):
    query: dict = {}
    if mine or user["role"] == "cashier":
        query["cashier"] = user["id"]
    if fromDate:
        query.setdefault("saleDate", {})["$gte"] = fromDate
    if toDate:
        query.setdefault("saleDate", {})["$lte"] = toDate + "T23:59:59Z"
    cur = db.sales.find(query, {"_id": 0}).sort("saleDate", -1).limit(limit)
    return await cur.to_list(limit)


@api.get("/sales/{sid}")
async def get_sale(sid: str, user: dict = Depends(get_current_user)):
    sale = await db.sales.find_one({"id": sid}, {"_id": 0})
    if not sale:
        raise HTTPException(404, "Sale not found")
    if user["role"] == "cashier" and sale["cashier"] != user["id"]:
        raise HTTPException(403, "Not authorized")
    items = await db.sale_items.find({"sale": sid}, {"_id": 0}).to_list(500)
    payments = await db.payments.find({"sale": sid}, {"_id": 0}).to_list(20)
    settings = await db.store_settings.find_one({"id": "default"}, {"_id": 0})
    return {"sale": sale, "items": items, "payments": payments, "store": settings or {}}


@api.post("/sales/{sid}/cancel")
async def cancel_sale(sid: str, admin: dict = Depends(require_admin)):
    sale = await db.sales.find_one({"id": sid})
    if not sale:
        raise HTTPException(404, "Sale not found")
    if sale["status"] == "cancelled":
        raise HTTPException(400, "Sale already cancelled")
    # restore stock
    items = await db.sale_items.find({"sale": sid}, {"_id": 0}).to_list(500)
    for si in items:
        p = await db.products.find_one({"id": si["product"]})
        if p:
            prev = float(p.get("stockQuantity", 0))
            new_stock = prev + float(si["quantity"])
            await db.products.update_one({"id": p["id"]}, {"$set": {"stockQuantity": new_stock}})
            await db.stock_movements.insert_one({
                "id": str(uuid.uuid4()),
                "product": p["id"],
                "productName": p["productName"],
                "movementType": "correction",
                "quantity": si["quantity"],
                "previousStock": prev,
                "newStock": new_stock,
                "reason": f"Cancelled sale {sale['invoiceNumber']}",
                "createdBy": admin["id"],
                "createdByName": admin["fullName"],
                "createdAt": now_iso(),
            })
    await db.sales.update_one(
        {"id": sid},
        {"$set": {"status": "cancelled", "cancelledBy": admin["id"], "cancelledAt": now_iso()}},
    )
    return {"success": True}


# ---------------- Dashboard ----------------
@api.get("/dashboard/stats")
async def dashboard_stats(_: dict = Depends(require_admin)):
    today_start = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0).isoformat()
    month_start = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0, microsecond=0).isoformat()

    today_sales = await db.sales.find(
        {"saleDate": {"$gte": today_start}, "status": "completed"}, {"_id": 0}
    ).to_list(10000)
    month_sales = await db.sales.find(
        {"saleDate": {"$gte": month_start}, "status": "completed"}, {"_id": 0}
    ).to_list(10000)

    today_total = sum(s["grandTotal"] for s in today_sales)
    month_total = sum(s["grandTotal"] for s in month_sales)

    products_count = await db.products.count_documents({"status": "active"})
    users_count = await db.users.count_documents({"role": "cashier", "status": "active"})

    low_stock_cursor = db.products.find({"status": "active"}, {"_id": 0})
    low_stock = []
    async for p in low_stock_cursor:
        if float(p.get("stockQuantity", 0)) <= float(p.get("minimumStockLevel", 0)):
            low_stock.append(p)
    low_stock = sorted(low_stock, key=lambda x: x.get("stockQuantity", 0))[:10]

    # Recent sales
    recent = await db.sales.find({}, {"_id": 0}).sort("saleDate", -1).limit(8).to_list(8)

    # 7-day chart
    chart = []
    for i in range(6, -1, -1):
        day = (datetime.now(timezone.utc) - timedelta(days=i)).replace(hour=0, minute=0, second=0, microsecond=0)
        next_day = day + timedelta(days=1)
        day_sales = await db.sales.find(
            {"saleDate": {"$gte": day.isoformat(), "$lt": next_day.isoformat()}, "status": "completed"},
            {"_id": 0, "grandTotal": 1},
        ).to_list(10000)
        chart.append({
            "date": day.strftime("%b %d"),
            "total": round(sum(s["grandTotal"] for s in day_sales), 2),
            "count": len(day_sales),
        })

    return {
        "todayTotal": round(today_total, 2),
        "todayCount": len(today_sales),
        "monthTotal": round(month_total, 2),
        "monthCount": len(month_sales),
        "productsCount": products_count,
        "cashiersCount": users_count,
        "lowStock": low_stock,
        "recentSales": recent,
        "chart": chart,
    }


# ---------------- Reports ----------------
@api.get("/reports/sales")
async def report_sales(
    _: dict = Depends(require_admin),
    fromDate: Optional[str] = None,
    toDate: Optional[str] = None,
):
    q = {"status": "completed"}
    if fromDate:
        q.setdefault("saleDate", {})["$gte"] = fromDate
    if toDate:
        q.setdefault("saleDate", {})["$lte"] = toDate + "T23:59:59Z"
    sales = await db.sales.find(q, {"_id": 0}).sort("saleDate", -1).to_list(10000)
    total = sum(s["grandTotal"] for s in sales)
    tax = sum(s.get("taxAmount", 0) for s in sales)
    discount = sum(s.get("discountAmount", 0) for s in sales)
    return {"sales": sales, "summary": {"total": round(total, 2), "tax": round(tax, 2), "discount": round(discount, 2), "count": len(sales)}}


async def _sale_ids_in_range(fromDate: Optional[str], toDate: Optional[str]) -> Optional[List[str]]:
    if not fromDate and not toDate:
        return None
    q: dict = {"status": "completed"}
    if fromDate:
        q.setdefault("saleDate", {})["$gte"] = fromDate
    if toDate:
        q.setdefault("saleDate", {})["$lte"] = toDate + "T23:59:59Z"
    docs = await db.sales.find(q, {"_id": 0, "id": 1}).to_list(100000)
    return [d["id"] for d in docs]


@api.get("/reports/top-products")
async def report_top_products(
    _: dict = Depends(require_admin),
    limit: int = 10,
    fromDate: Optional[str] = None,
    toDate: Optional[str] = None,
):
    match: dict = {}
    sale_ids = await _sale_ids_in_range(fromDate, toDate)
    if sale_ids is not None:
        match = {"sale": {"$in": sale_ids}}
    pipeline = [
        {"$match": match} if match else {"$match": {}},
        {"$group": {"_id": "$product", "productName": {"$first": "$productName"}, "qty": {"$sum": "$quantity"}, "revenue": {"$sum": "$totalAmount"}}},
        {"$sort": {"qty": -1}},
        {"$limit": limit},
    ]
    res = await db.sale_items.aggregate(pipeline).to_list(limit)
    return [{"productId": r["_id"], "productName": r["productName"], "qty": r["qty"], "revenue": round(r["revenue"], 2)} for r in res]


@api.get("/reports/slow-products")
async def report_slow_products(
    _: dict = Depends(require_admin),
    limit: int = 10,
    fromDate: Optional[str] = None,
    toDate: Optional[str] = None,
):
    match: dict = {}
    sale_ids = await _sale_ids_in_range(fromDate, toDate)
    if sale_ids is not None:
        match = {"sale": {"$in": sale_ids}}
    pipeline = [
        {"$match": match} if match else {"$match": {}},
        {"$group": {"_id": "$product", "productName": {"$first": "$productName"}, "qty": {"$sum": "$quantity"}}},
        {"$sort": {"qty": 1}},
        {"$limit": limit},
    ]
    return [
        {"productId": r["_id"], "productName": r["productName"], "qty": r["qty"]}
        for r in await db.sale_items.aggregate(pipeline).to_list(limit)
    ]


@api.get("/reports/cashiers")
async def report_cashiers(
    _: dict = Depends(require_admin),
    fromDate: Optional[str] = None,
    toDate: Optional[str] = None,
):
    q: dict = {"status": "completed"}
    if fromDate:
        q.setdefault("saleDate", {})["$gte"] = fromDate
    if toDate:
        q.setdefault("saleDate", {})["$lte"] = toDate + "T23:59:59Z"
    sales = await db.sales.find(q, {"_id": 0}).to_list(100000)
    by_cashier: dict = {}
    for s in sales:
        cid = s["cashier"]
        if cid not in by_cashier:
            by_cashier[cid] = {
                "cashierId": cid,
                "cashierName": s.get("cashierName", "—"),
                "orders": 0,
                "revenue": 0.0,
                "discount": 0.0,
                "tax": 0.0,
                "saleIds": [],
            }
        b = by_cashier[cid]
        b["orders"] += 1
        b["revenue"] += s["grandTotal"]
        b["discount"] += s.get("discountAmount", 0)
        b["tax"] += s.get("taxAmount", 0)
        b["saleIds"].append(s["id"])

    # Top products per cashier
    result = []
    for c in by_cashier.values():
        items_cursor = db.sale_items.aggregate([
            {"$match": {"sale": {"$in": c["saleIds"]}}},
            {"$group": {"_id": "$product", "productName": {"$first": "$productName"}, "qty": {"$sum": "$quantity"}, "revenue": {"$sum": "$totalAmount"}}},
            {"$sort": {"qty": -1}},
            {"$limit": 50},
        ])
        items = await items_cursor.to_list(50)
        result.append({
            "cashierId": c["cashierId"],
            "cashierName": c["cashierName"],
            "orders": c["orders"],
            "revenue": round(c["revenue"], 2),
            "discount": round(c["discount"], 2),
            "tax": round(c["tax"], 2),
            "products": [
                {"productId": i["_id"], "productName": i["productName"], "qty": i["qty"], "revenue": round(i["revenue"], 2)}
                for i in items
            ],
        })
    result.sort(key=lambda x: x["revenue"], reverse=True)
    return result


@api.get("/reports/stock")
async def report_stock(_: dict = Depends(require_admin)):
    products = await db.products.find({}, {"_id": 0}).sort("productName", 1).to_list(10000)
    valuation = sum(float(p.get("costPrice", 0)) * float(p.get("stockQuantity", 0)) for p in products)
    selling_valuation = sum(float(p.get("sellingPrice", 0)) * float(p.get("stockQuantity", 0)) for p in products)
    return {
        "products": products,
        "summary": {
            "totalProducts": len(products),
            "costValuation": round(valuation, 2),
            "sellingValuation": round(selling_valuation, 2),
        },
    }


# ---------------- Store Settings ----------------
@api.get("/settings")
async def get_settings(user: dict = Depends(get_current_user)):
    doc = await db.store_settings.find_one({"id": "default"}, {"_id": 0})
    if not doc:
        doc = {
            "id": "default",
            "storeName": "My Store",
            "storeAddress": "",
            "mobileNumber": "",
            "email": "",
            "gstNumber": "",
            "taxEnabled": True,
            "invoicePrefix": "INV",
            "receiptFooter": "Thank you for shopping with us!",
        }
    return doc


@api.put("/settings")
async def update_settings(payload: StoreSettingsIn, _: dict = Depends(require_admin)):
    data = {"id": "default", **payload.model_dump(), "updatedAt": now_iso()}
    await db.store_settings.update_one({"id": "default"}, {"$set": data}, upsert=True)
    data.pop("_id", None)
    return data


# ---------------- Health ----------------
@api.get("/")
async def root():
    return {"name": "POS Billing API", "status": "ok"}


# ---------------- Seeding & Indexes ----------------
async def ensure_indexes():
    await db.users.create_index("username", unique=True)
    await db.products.create_index("productName")
    await db.products.create_index("barcode", sparse=True)
    await db.products.create_index("sku", sparse=True)
    await db.products.create_index("category")
    await db.sales.create_index("invoiceNumber", unique=True)
    await db.sales.create_index("saleDate")
    await db.sales.create_index("cashier")
    await db.sale_items.create_index("sale")
    await db.stock_movements.create_index("product")


async def seed_data():
    if not await db.users.find_one({"username": "admin"}):
        await db.users.insert_one({
            "id": str(uuid.uuid4()),
            "username": "admin",
            "fullName": "Store Admin",
            "mobileNumber": "9999999999",
            "passwordHash": hash_pw("admin123"),
            "role": "admin",
            "status": "active",
            "createdAt": now_iso(),
            "updatedAt": now_iso(),
        })
        logger.info("Seeded admin user (admin/admin123)")

    if not await db.users.find_one({"username": "cashier"}):
        await db.users.insert_one({
            "id": str(uuid.uuid4()),
            "username": "cashier",
            "fullName": "Demo Cashier",
            "mobileNumber": "8888888888",
            "passwordHash": hash_pw("cashier123"),
            "role": "cashier",
            "status": "active",
            "createdAt": now_iso(),
            "updatedAt": now_iso(),
        })
        logger.info("Seeded cashier user (cashier/cashier123)")

    if not await db.store_settings.find_one({"id": "default"}):
        await db.store_settings.insert_one({
            "id": "default",
            "storeName": "Sunrise Mart",
            "storeAddress": "123 Market Street, Mumbai, India",
            "mobileNumber": "+91 98765 43210",
            "email": "hello@sunrisemart.com",
            "gstNumber": "27AABCS1234A1Z5",
            "taxEnabled": True,
            "invoicePrefix": "INV",
            "receiptFooter": "Thank you for shopping with Sunrise Mart!",
            "createdAt": now_iso(),
            "updatedAt": now_iso(),
        })

    if await db.categories.count_documents({}) == 0:
        cats = [
            ("Beverages", "Coffee, tea, juices, soft drinks"),
            ("Bakery", "Bread, croissants, pastries"),
            ("Groceries", "Daily essentials"),
            ("Snacks", "Chips, biscuits, chocolates"),
            ("Dairy", "Milk, cheese, butter, yoghurt"),
        ]
        cat_ids = {}
        for name, desc in cats:
            cid = str(uuid.uuid4())
            cat_ids[name] = cid
            await db.categories.insert_one({
                "id": cid, "categoryName": name, "description": desc,
                "status": "active", "createdAt": now_iso(), "updatedAt": now_iso(),
            })

        # Sample supplier
        sid = str(uuid.uuid4())
        await db.suppliers.insert_one({
            "id": sid, "supplierName": "Global Distributors",
            "contactPerson": "Ravi Kumar", "mobileNumber": "+91 99887 76655",
            "email": "ravi@globaldist.com", "address": "Andheri East, Mumbai",
            "status": "active", "createdAt": now_iso(), "updatedAt": now_iso(),
        })

        products = [
            ("Cappuccino", "Beverages", 120, 180, 5, 50, 10),
            ("Espresso Shot", "Beverages", 60, 110, 5, 100, 20),
            ("Cold Brew", "Beverages", 140, 220, 5, 30, 8),
            ("Butter Croissant", "Bakery", 40, 90, 5, 25, 10),
            ("Almond Croissant", "Bakery", 55, 120, 5, 4, 10),
            ("Sourdough Loaf", "Bakery", 90, 180, 5, 12, 5),
            ("Rice Bag 5kg", "Groceries", 350, 480, 0, 40, 10),
            ("Sunflower Oil 1L", "Groceries", 140, 180, 5, 60, 15),
            ("Whole Wheat Flour 1kg", "Groceries", 50, 75, 0, 80, 20),
            ("Potato Chips Salted", "Snacks", 25, 50, 12, 200, 30),
            ("Chocolate Cookies", "Snacks", 30, 60, 12, 150, 30),
            ("Dark Chocolate Bar", "Snacks", 80, 150, 12, 5, 10),
            ("Fresh Milk 1L", "Dairy", 45, 65, 0, 70, 15),
            ("Greek Yoghurt 500g", "Dairy", 90, 140, 5, 35, 10),
            ("Cheddar Cheese 200g", "Dairy", 180, 280, 5, 22, 5),
        ]
        for name, cat, cost, sell, tax, stock, minimum in products:
            await db.products.insert_one({
                "id": str(uuid.uuid4()),
                "productName": name,
                "barcode": str(8900000000000 + abs(hash(name)) % 999999),
                "sku": name.replace(" ", "-").upper()[:18],
                "category": cat_ids[cat],
                "brand": "House",
                "costPrice": cost,
                "sellingPrice": sell,
                "taxPercentage": tax,
                "stockQuantity": stock,
                "minimumStockLevel": minimum,
                "supplier": sid,
                "productImageUrl": "",
                "status": "active",
                "createdAt": now_iso(),
                "updatedAt": now_iso(),
            })
        logger.info("Seeded categories, supplier, and %d products", len(products))


@app.on_event("startup")
async def on_startup():
    try:
        await ensure_indexes()
        await seed_data()
    except Exception as e:
        logger.exception("Startup error: %s", e)


@app.on_event("shutdown")
async def on_shutdown():
    client.close()


app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8001))
    uvicorn.run("server:app", host="0.0.0.0", port=port, reload=True)

